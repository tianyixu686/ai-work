"""Build the course report and supporting figures."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "code"))
from expert_system import recognize  # noqa: E402

ACCENT = "315B7D"
LIGHT = "EAF1F6"
GRID = "D9D9D9"
FONT = "/System/Library/Fonts/STHeiti Medium.ttc"


def font(size, bold=False):
    return ImageFont.truetype(FONT, size=size, index=0)


def flow_figure(out):
    img = Image.new("RGB", (1600, 620), "white")
    d = ImageDraw.Draw(img)
    title = font(42); body = font(31); small = font(27)
    d.text((70, 35), "两种动物识别路线", fill="#111111", font=title)
    boxes = [
        (90, 170, 350, 320, "输入图像", "RGB 32×32"),
        (470, 95, 840, 245, "产生式系统", "统计特征 → 事实 → 规则链"),
        (470, 335, 840, 485, "轻量 CNN", "卷积特征 → 全局池化"),
        (990, 170, 1440, 320, "四类输出", "猫  狗  鹿  马"),
    ]
    for x1, y1, x2, y2, t, s in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=20, fill="#EAF1F6", outline="#315B7D", width=4)
        d.text(((x1+x2)//2, y1+35), t, fill="#111111", font=body, anchor="ma")
        d.text(((x1+x2)//2, y1+92), s, fill="#333333", font=small, anchor="ma")
    for a, b in [((350,245),(470,170)),((350,245),(470,410)),((840,170),(990,245)),((840,410),(990,245))]:
        d.line((*a,*b), fill="#315B7D", width=6)
        x,y=b; d.polygon([(x,y),(x-22,y-12),(x-22,y+12)], fill="#315B7D")
    img.save(out)


def network_figure(out):
    img = Image.new("RGB", (1600, 520), "white"); d = ImageDraw.Draw(img)
    labels = [("3×32×32", "输入"), ("32×32×32", "3×3卷积"), ("48×16×16", "DS残差块"),
              ("96×8×8", "DS残差块"), ("160×4×4", "DS残差块"), ("160", "全局池化"), ("4", "Softmax")]
    x = 45
    for i, (shape, label) in enumerate(labels):
        w = 180 if i < 5 else 165
        d.rounded_rectangle((x,150,x+w,330), radius=16, fill="#EAF1F6" if i%2==0 else "#F5F7F9", outline="#315B7D", width=4)
        d.text((x+w/2,205), shape, fill="#111", font=font(28), anchor="mm")
        d.text((x+w/2,270), label, fill="#333", font=font(24), anchor="mm")
        if i < len(labels)-1:
            d.line((x+w,240,x+w+42,240), fill="#315B7D", width=5)
            d.polygon([(x+w+42,240),(x+w+26,230),(x+w+26,250)], fill="#315B7D")
        x += w + 45
    d.text((50,45), "AnimalLiteNet 结构", fill="#111", font=font(40))
    d.text((50,410), "DS 表示深度可分离卷积；残差支路改善梯度传播。", fill="#333", font=font(28))
    img.save(out)


def examples_figure(out):
    canvas = Image.new("RGB", (1480, 420), "white"); d = ImageDraw.Draw(canvas)
    names = [("cat", "猫"), ("dog", "狗"), ("deer", "鹿"), ("horse", "马")]
    for i, (fn, cn) in enumerate(names):
        im = Image.open(ROOT/"examples"/f"{fn}.png").convert("RGB").resize((280,280), Image.Resampling.NEAREST)
        x = 70 + i*355; canvas.paste(im,(x,35))
        d.rectangle((x,35,x+280,315), outline="#D9D9D9", width=3)
        d.text((x+140,365), cn, fill="#111", font=font(32), anchor="mm")
    canvas.save(out)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = OxmlElement(f"w:{m}"); node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa"); tcMar.append(node)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
    hdr = table.rows[0]; set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]; c.text=h; set_cell_shading(c,ACCENT); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(10)
        c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; set_cell_margins(c)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            if ri%2: set_cell_shading(cells[i],"F3F7FA")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.size=Pt(9.5)
                if i==0 and len(str(val))<10: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return table


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_code(doc, code):
    p=doc.add_paragraph(); p.style="No Spacing"; p.paragraph_format.left_indent=Cm(.6); p.paragraph_format.right_indent=Cm(.4)
    for line in code.strip().splitlines():
        r=p.add_run(line+"\n"); r.font.name="Menlo"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"Menlo"); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(45,45,45)


def build():
    figures=ROOT/"figures"; figures.mkdir(exist_ok=True)
    flow_figure(figures/"method_flow.png"); network_figure(figures/"network.png"); examples_figure(figures/"examples.png")
    demo={n:recognize(ROOT/"examples"/f"{n}.png") for n in ("cat","dog","deer","horse")}
    (ROOT/"results"/"rule_demo.json").write_text(json.dumps(demo,ensure_ascii=False,indent=2),encoding="utf-8")

    doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.3); sec.right_margin=Cm(2.3)
    styles=doc.styles
    normal=styles["Normal"]; normal.font.name="STHeiti"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"STHeiti"); normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.35; normal.paragraph_format.space_after=Pt(6)
    for name,size,bold in [("Title",22,True),("Heading 1",16,True),("Heading 2",13,True),("Heading 3",11.5,True)]:
        s=styles[name]; s.font.name="STHeiti"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"STHeiti"); s.font.size=Pt(size); s.font.bold=bold; s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_before=Pt(12 if name!="Title" else 0); s.paragraph_format.space_after=Pt(6)
    styles["Title"].paragraph_format.space_after=Pt(18)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    for child in list(title_ppr):
        if child.tag == qn("w:pBdr"):
            title_ppr.remove(child)

    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("基于产生式系统与轻量 CNN 的动物识别方法比较")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("课程作业报告"); r.font.size=Pt(14); r.font.bold=True
    doc.add_paragraph("\n\n")
    meta=doc.add_table(rows=4,cols=2); meta.alignment=1
    for i,(a,b) in enumerate([("\u59d3\u540d","\u5f90\u5929\u4e00"),("\u8bc6\u522b\u5bf9\u8c61","\u732b\u3001\u72d7\u3001\u9e7f\u3001\u9a6c"),("\u5f00\u53d1\u8bed\u8a00","Python"),("\u6846\u67b6","PyTorch")]):
        meta.cell(i,0).text=a; meta.cell(i,1).text=b
        for c in meta.rows[i].cells: set_cell_margins(c,180,200,180,200); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(meta.cell(i,0),LIGHT)
    doc.add_paragraph("\n")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("本报告实现一套可解释的 IF–THEN 识别器与一个端到端轻量 CNN，并从原理、流程、性能和应用边界进行比较。")
    doc.add_page_break()

    add_heading(doc,"摘要",1)
    doc.add_paragraph("动物图像在光照、姿态、遮挡和背景上差异明显。本作业围绕猫、狗、鹿和马四类目标，实现两种识别方法。第一种将颜色、边缘、纹理、中心与边界亮度差转化为符号事实，通过产生式规则正向链推理；第二种使用深度可分离卷积和残差连接学习图像特征。两者的核心差异是：产生式系统依赖人工选择特征与阈值，可以输出清晰推理链；CNN 通过数据自动获得多层表征，通常具有更好的复杂场景适应性，但需要训练数据和计算资源。")
    doc.add_paragraph("关键词：产生式系统；动物识别；卷积神经网络；深度可分离卷积；可解释性")
    add_heading(doc,"一 问题定义与总体方案",1)
    doc.add_paragraph("任务输入为一张 RGB 动物图像，输出为四类中的一类。数据采用 CIFAR-10 中的 cat、dog、deer 和 horse 子集。为了使比较具有可重复性，代码固定随机种子 2026，从原始训练集中划分 85% 训练数据和 15% 验证数据，原始测试集仅用于最终评估。")
    doc.add_picture(str(figures/"method_flow.png"),width=Inches(6.45)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph("图 1  两种识别方法的总体流程"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc,"二 产生式动物识别方法",1)
    add_heading(doc,"二点一 基本原理",2)
    doc.add_paragraph("产生式系统由规则库、工作记忆和推理机组成。每条规则写成「IF 条件集 THEN 结论」。系统先从图像中计算七个统计量，再根据阈值生成 warm_color、rough_texture 等事实。正向链从已知事实出发，反复触发所有前件已满足的规则，直到不再出现新事实。")
    add_table(doc,["统计量","计算方式","用途"],[
        ("RGB 均值","各通道像素均值","判断暖色、冷色及蓝色背景"),
        ("边缘强度","水平与垂直灰度差的均值之和","近似表示轮廓与细节多少"),
        ("纹理强度","灰度标准差","区分平滑与粗糙外观"),
        ("中心对比","中心区均值减边界区均值","判断主体是否亮于背景"),
        ("方向比","水平梯度与垂直梯度之比","粗略表示水平延展或细长结构"),
    ],[3.0,7.0,5.5])
    add_heading(doc,"二点二 规则设计",2)
    add_table(doc,["规则","条件","结论"],[
        ("R1","smooth_texture 且 horizontal_body","candidate_cat"),
        ("R2","warm_color 且 slender_shape","candidate_dog"),
        ("R3","blue_background 且 low_center_contrast","candidate_deer"),
        ("R4","rough_texture 且 horizontal_body","candidate_horse"),
        ("R5–R8","候选类别与二级特征同时满足","猫、狗、鹿或马"),
    ],[2.0,8.5,5.0])
    doc.add_paragraph("规则的两层结构使「候选类别」和「最终类别」分离，方便查看推理链。若没有规则得到最终类别，系统返回 unknown，而不强制选择一类。这是它与普通 Softmax 分类器的一个重要差异。")
    add_heading(doc,"二点三 核心代码",2)
    add_code(doc,"""
while changed:
    changed = False
    for rule in RULES:
        if rule.conditions <= memory and rule.conclusion not in memory:
            memory.add(rule.conclusion)
            trace.append({"rule": rule.name, "then": rule.conclusion})
            changed = True
""")
    doc.add_paragraph("代码使用集合包含关系检查前件，每次触发都保存规则名和结论，因此可以直接输出推理过程。")

    add_heading(doc,"三 基于轻量 CNN 的动物识别",1)
    add_heading(doc,"三点一 模型结构",2)
    doc.add_paragraph("卷积神经网络通过可学习卷积核从局部像素模式中逐级构建边缘、纹理、部件和类别表征。本作业不使用 VGG 那样的大量普通卷积，而采用「逐通道 3×3 卷积 + 1×1 逐点卷积」。这种深度可分离卷积能降低参数量和计算量。残差捷径把输入与卷积分支相加，减少深层网络的优化困难。")
    doc.add_picture(str(figures/"network.png"),width=Inches(6.45)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph("图 2  轻量 CNN 的层级结构"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc,"三点二 数据预处理",2)
    doc.add_paragraph("训练阶段采用随机填充裁剪、水平翻转和轻度颜色抖动，提高对位置、朝向和光照变化的适应性。验证和测试阶段不使用随机增强，只执行归一化，以保证结果稳定。")
    doc.add_picture(str(figures/"examples.png"),width=Inches(6.45)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph("图 3  四类低分辨率动物图像示例"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc,"三点三 训练策略",2)
    add_table(doc,["项目","设置","原因"],[
        ("损失函数","CrossEntropyLoss，label smoothing=0.05","抑制过度自信"),
        ("优化器","AdamW，lr=0.002，weight decay=1e-4","训练稳定且具有权重正则化"),
        ("学习率","CosineAnnealingLR","后期平滑降低学习率"),
        ("模型选择","验证准确率最高的权重","测试集不参与选择"),
        ("早停","连续 5 轮无改善","减少过拟合和无效计算"),
    ],[3,5.5,7])

    add_heading(doc,"四 实验方法与结果记录",1)
    add_heading(doc,"四点一 评价指标",2)
    doc.add_paragraph("分类结果建议同时报告总体准确率、各类召回率、宏平均 F1 值和混淆矩阵。准确率反映整体正确比例；召回率表示某类真实样本被找回的比例；宏 F1 对每个类别等权处理，不会被大类别主导。")
    add_heading(doc,"四点二 可复现结果的记录方式",2)
    doc.add_paragraph("为避免把未实际运行的数值写入报告，本提交不伪造 CNN 准确率。运行 `python code/train.py` 后，脚本会自动生成 `results/history.csv`、`results/metrics.json` 和 `results/best_model.pt`。评估时应将下表的「运行后填入」替换为 metrics.json 中的实际值。")
    add_table(doc,["方法","准确率","可解释输出","训练需求","局限"],[
        ("产生式系统","需在完整测试集评估","事实与规则链","无梯度训练","阈值对背景敏感"),
        ("AnimalLiteNet","metrics.json 运行后填入","类别概率","需要数据与迭代训练","决策过程不直观"),
    ],[3,3.2,3.6,3.4,4.2])
    add_heading(doc,"四点三 附带样例的规则演示",2)
    rows=[]; zh={"cat":"猫","dog":"狗","deer":"鹿","horse":"马",None:"未知"}
    for name in ("cat","dog","deer","horse"):
        r=demo[name]; rows.append((zh[name],zh[r["prediction"]],"→".join(x["rule"] for x in r["trace"]) or "无规则命中","正确" if r["prediction"]==name else "错误/拒识"))
    add_table(doc,["样例","规则结果","触发规则","判定"],rows,[2.3,3.0,6.3,3.3])
    doc.add_paragraph("这四张图只是代码功能演示，不能代替完整测试集上的统计评估。产生式系统的优势是每个结论都可以追溯到事实和规则；同时，单张图的统计阈值容易受背景影响，因此必须用大量样本验证。")

    add_heading(doc,"五 两种方法的对比分析",1)
    add_table(doc,["比较维度","产生式系统","轻量 CNN"],[
        ("知识来源","人工设计特征、阈值和规则","从标注数据中自动学习参数"),
        ("表征能力","主要是低阶统计特征","可形成从边缘到语义部件的层级表征"),
        ("可解释性","高，可显示完整规则链","中等或较低，通常需要 Grad-CAM 等辅助工具"),
        ("数据与计算","数据要求低，推理快","训练需较多数据和算力，训练后推理较快"),
        ("维护方式","增删规则与调整阈值","补充数据、调参并重新训练"),
        ("复杂场景","容易出现规则爆炸和冲突","对姿态、背景和纹理变化更有弹性"),
    ],[3.0,6.2,6.2])
    doc.add_paragraph("两种方法不是简单的「旧方法与新方法」关系。当规则可被专家准确表达、数据稀少或需要审核每个决策时，产生式系统仍具有价值。当图像变化较大且有足够数据时，CNN 更合适。在实际系统中，可以让 CNN 负责图像感知，再用规则系统表达业务约束、风险判断和拒识逻辑。")

    add_heading(doc,"六 实现步骤与运行说明",1)
    for text in [
        "安装 Python 3.10 或更高版本，在作业目录中创建虚拟环境并执行 pip install -r requirements.txt。",
        "运行 python code/expert_system.py examples/cat.png，查看图像统计、初始事实、规则链和结论。",
        "运行 python code/train.py --epochs 20 --batch-size 128。首次会下载 CIFAR-10，最优权重保存为 results/best_model.pt。",
        "运行 python code/predict.py examples/dog.png --weights results/best_model.pt，对同一张图比较规则结论和 CNN 概率。",
        "用完整测试集统计混淆矩阵和分类别召回率，并把实际环境、随机种子和软件版本一并记录。",
    ]:
        p=doc.add_paragraph(style="List Number"); p.add_run(text)

    add_heading(doc,"七 局限与改进方向",1)
    doc.add_paragraph("第一，CIFAR-10 图像只有 32×32 像素，与手机拍摄或监控图像存在显著域差异；第二，四类封闭集设定不能处理未知动物，CNN 即使面对集外图像也会给出四类概率；第三，目前规则使用全图统计量，未先定位动物主体。可以通过目标检测后再分类、引入开放集拒识、扩充真实场景数据、使用预训练模型以及进行温度缩放等置信度校准来改进。")
    add_heading(doc,"八 结论",1)
    doc.add_paragraph("本作业完成了一套可运行的产生式动物识别器与一个轻量 CNN 分类器。产生式系统展示了「特征—事实—规则—结论」的显式推理，便于理解和修改；CNN 通过深度可分离卷积学习更灵活的图像表征。方法选择应由数据规模、计算条件、准确性目标和解释性要求共同决定。")
    add_heading(doc,"参考文献",1)
    refs=[
        "[1] Krizhevsky A. Learning Multiple Layers of Features from Tiny Images. 2009.",
        "[2] LeCun Y, Bengio Y, Hinton G. Deep learning. Nature, 2015, 521: 436–444.",
        "[3] He K, Zhang X, Ren S, Sun J. Deep Residual Learning for Image Recognition. CVPR, 2016.",
        "[4] Howard A G, et al. MobileNets: Efficient Convolutional Neural Networks for Mobile Vision Applications. arXiv:1704.04861, 2017.",
        "[5] PyTorch Documentation. https://pytorch.org/docs/stable/",
        "[6] CIFAR-10 Dataset. https://www.cs.toronto.edu/~kriz/cifar.html",
    ]
    for r in refs: doc.add_paragraph(r)

    for section in doc.sections:
        footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=footer.add_run("徐天一  ·  动物识别课程作业  ·  ")
        fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); run._r.addnext(fld)
    out=ROOT/"动物识别+徐天一.docx"; doc.save(out); print(out)


if __name__ == "__main__": build()

